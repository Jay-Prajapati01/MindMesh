"""Document ingestion for DOCX, PPTX, TXT, Markdown, CSV, and Excel.

Each format has a dedicated extractor that returns plain text. The public
`extract_text` dispatches on the file extension.
"""
from __future__ import annotations

import csv
import io

from app.core.logging import get_logger

logger = get_logger(__name__)

SUPPORTED = {".docx", ".pptx", ".txt", ".md", ".markdown", ".csv", ".xlsx"}


def _extract_docx(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        slide_parts = [f"Slide {i}:"]
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                slide_parts.append(shape.text_frame.text.strip())
        if len(slide_parts) > 1:
            parts.append("\n".join(slide_parts))
    return "\n\n".join(parts).strip()


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts).strip()


def _extract_csv(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = [" | ".join(row) for row in reader if any(cell.strip() for cell in row)]
    return "\n".join(rows).strip()


def _extract_plain(data: bytes) -> str:
    return data.decode("utf-8", errors="replace").strip()


def extract_text(data: bytes, filename: str) -> dict:
    """Extract text from a supported document. Returns {title, text}."""
    lower = filename.lower()
    ext = "." + lower.rsplit(".", 1)[-1] if "." in lower else ""
    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported document type: {ext or filename}")

    if ext == ".docx":
        text = _extract_docx(data)
    elif ext == ".pptx":
        text = _extract_pptx(data)
    elif ext == ".xlsx":
        text = _extract_xlsx(data)
    elif ext == ".csv":
        text = _extract_csv(data)
    else:  # .txt, .md, .markdown
        text = _extract_plain(data)

    return {"title": filename, "text": text}
